"""Prep a real Greenhouse posting + dummy docs for an apply.py dry-run.

Inserts the maintainx Acquisition Lead posting via the tracker
(source='manual_entry'), then creates placeholder resume + cover
letter .docx files in data/{profile}/applications/pending/ so the
agent has something to upload during the dry-run.

The .docx contents do NOT matter for selector testing -- the
dry-run never submits, so the placeholder data never reaches the
employer. Swap with real Phase 10 output before any real submission.
"""
from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from docx import Document
from engine.persistence.tracker import Tracker


POSTING = {
    "company": "MaintainX",
    "url": (
        "https://job-boards.greenhouse.io/maintainx/jobs/5120026007"
    ),
    "title": "Acquisition Lead",
    "location": (
        "Toronto, ON / Montreal / San Francisco / Miami / Raleigh"
    ),
    "posting_text": (
        "MaintainX is hiring an Acquisition Lead. "
        "Placeholder body for dry-run testing of the Greenhouse "
        "ATS handler. Sourced via manual_entry seed."
    ),
}


def _make_dummy_docx(path: Path, header: str, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading(header, level=1)
    doc.add_paragraph(body)
    doc.save(str(path))


def main() -> int:
    profile = "default"
    t = Tracker(profile)
    try:
        cid = t.upsert_company(POSTING["company"])
        opp_id, was_new = t.insert_opportunity(
            company_id=cid,
            source="manual_entry",
            source_url=POSTING["url"],
            title=POSTING["title"],
            location=POSTING["location"],
            posting_text=POSTING["posting_text"],
        )
        print(f"Posting #{opp_id}  (was_new={was_new})")
    finally:
        t.close()

    pending = (
        PROJECT_ROOT / "data" / profile / "applications" / "pending"
    )
    resume = pending / f"resume_{opp_id}.docx"
    cl = pending / f"cover_letter_{opp_id}.docx"
    _make_dummy_docx(
        resume,
        "Alex Doe",
        "Placeholder resume body for dry-run. "
        "Swap with real Phase 10 output before any real submission.",
    )
    _make_dummy_docx(
        cl, "Cover Letter",
        "Placeholder cover letter body for dry-run.",
    )
    print(f"Resume:       {resume}")
    print(f"Cover letter: {cl}")
    print()
    print(
        f"Run: python scripts/apply.py --profile default "
        f"--posting {opp_id} --dry-run"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
