"""Step 3: render Claude's markdown output into an ATS-friendly
.docx file.

Parses the markdown structure produced by ResumePromptBuilder's
OUTPUT_MARKDOWN_TEMPLATE (or CoverLetterPromptBuilder's) and emits
a single-column Word document that satisfies the canonical ATS
formatting rules:

  - .docx, single column, no tables / text boxes / columns / images
  - Body content only — NOTHING in the header/footer
  - Calibri font: 11pt body, 13pt section headings, 16pt name
  - 0.75 inch margins, 1.0–1.15 line spacing
  - Standard round bullets only
  - Standard section names (PROFESSIONAL SUMMARY, SKILLS,
    WORK EXPERIENCE, EDUCATION & CERTIFICATIONS)
  - Section order: Name+Contact → Summary → Skills → Experience →
    Education
  - File name: {FirstName}{LastName}Resume{CompanyName}{YYYYMMDD}.docx

The renderer is intentionally strict: it does not try to "fix"
malformed input. If the markdown doesn't follow the template, the
output may not parse cleanly downstream — fail visibly rather than
silently rearranging.
"""
from __future__ import annotations

import logging
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

logger = logging.getLogger(__name__)


# Canonical ATS sizes. Single source — tests pin to these.
NAME_FONT_SIZE_PT = 16
HEADING_FONT_SIZE_PT = 13
BODY_FONT_SIZE_PT = 11
FONT_NAME = "Calibri"
MARGIN_INCHES = 0.75
LINE_SPACING = 1.15


def _set_run_font(run, size_pt: int, bold: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold


def _set_section_margins(doc: Document) -> None:
    for section in doc.sections:
        section.top_margin = Inches(MARGIN_INCHES)
        section.bottom_margin = Inches(MARGIN_INCHES)
        section.left_margin = Inches(MARGIN_INCHES)
        section.right_margin = Inches(MARGIN_INCHES)


def _name_paragraph(doc: Document, name: str):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = LINE_SPACING
    run = para.add_run(name.strip())
    _set_run_font(run, NAME_FONT_SIZE_PT, bold=True)
    return para


def _centered_body_paragraph(doc: Document, text: str):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing = LINE_SPACING
    run = para.add_run(text.strip())
    _set_run_font(run, BODY_FONT_SIZE_PT)
    return para


def _heading_paragraph(doc: Document, heading: str):
    """Section heading like 'PROFESSIONAL SUMMARY'."""
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = LINE_SPACING
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(heading.upper().strip())
    _set_run_font(run, HEADING_FONT_SIZE_PT, bold=True)
    return para


def _body_paragraph(doc: Document, text: str):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = LINE_SPACING
    run = para.add_run(text)
    _set_run_font(run, BODY_FONT_SIZE_PT)
    return para


def _bullet_paragraph(doc: Document, text: str):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.line_spacing = LINE_SPACING
    run = para.add_run(text)
    _set_run_font(run, BODY_FONT_SIZE_PT)
    return para


def _skill_category_paragraph(doc: Document, category: str, terms: str):
    """`**Delivery & Program Management:** Agile, Scrum, ...`"""
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = LINE_SPACING
    bold_run = para.add_run(f"{category}: ")
    _set_run_font(bold_run, BODY_FONT_SIZE_PT, bold=True)
    norm_run = para.add_run(terms.strip())
    _set_run_font(norm_run, BODY_FONT_SIZE_PT)
    return para


def _role_header_paragraph(doc: Document, header_line: str):
    """Bold the title (everything before the first '|'), normal
    weight for company / location / dates after."""
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing = LINE_SPACING
    para.paragraph_format.space_before = Pt(6)
    if "|" in header_line:
        title, rest = header_line.split("|", 1)
        bold_run = para.add_run(title.strip())
        _set_run_font(bold_run, BODY_FONT_SIZE_PT, bold=True)
        normal_run = para.add_run(" | " + rest.strip())
        _set_run_font(normal_run, BODY_FONT_SIZE_PT)
    else:
        run = para.add_run(header_line)
        _set_run_font(run, BODY_FONT_SIZE_PT, bold=True)
    return para


# --- Markdown parsing --------------------------------------------

_SKILL_CATEGORY_RE = re.compile(r"^\*\*(.+?):\*\*\s*(.*)$")


def _parse_resume_markdown(content: str) -> dict:
    """Parse the resume markdown into a structured dict the renderer
    can walk.

    Returns:
      {
        "name": str,
        "tagline": str | None,
        "contact": str | None,
        "sections": [
          {"heading": "PROFESSIONAL SUMMARY", "blocks": [...]},
          ...
        ]
      }

    blocks are dicts of {"type": "paragraph"|"bullet"|"role_header"|
    "skill_category", "text": str | tuple}.

    The parser is forgiving about blank lines but strict about
    section markers (## HEADING) and role markers (### TITLE | ...).
    """
    lines = content.splitlines()
    name = ""
    tagline = None
    contact = None
    sections: list[dict] = []
    current_section: Optional[dict] = None
    pending_paragraph_lines: list[str] = []
    seen_h1 = False

    def _flush_pending():
        nonlocal pending_paragraph_lines
        if pending_paragraph_lines and current_section is not None:
            joined = " ".join(
                line.strip() for line in pending_paragraph_lines
                if line.strip()
            )
            if joined:
                current_section["blocks"].append(
                    {"type": "paragraph", "text": joined}
                )
        pending_paragraph_lines = []

    pre_section_lines: list[str] = []  # tagline + contact before first ##

    for raw in lines:
        line = raw.rstrip()
        # H1 — name (only first H1 honored)
        if line.startswith("# ") and not seen_h1:
            name = line[2:].strip()
            seen_h1 = True
            continue
        # Section heading
        if line.startswith("## "):
            _flush_pending()
            heading = line[3:].strip()
            current_section = {"heading": heading, "blocks": []}
            sections.append(current_section)
            # Anything we accumulated before the first ## is the
            # contact / tagline area
            if pre_section_lines:
                non_empty = [
                    p.strip() for p in pre_section_lines if p.strip()
                ]
                if non_empty:
                    tagline = non_empty[0] if len(non_empty) >= 1 else None
                    contact = non_empty[1] if len(non_empty) >= 2 else None
                pre_section_lines = []
            continue
        # Role header
        if line.startswith("### "):
            _flush_pending()
            if current_section is not None:
                current_section["blocks"].append({
                    "type": "role_header",
                    "text": line[4:].strip(),
                })
            continue
        # Bullet
        if line.lstrip().startswith("- "):
            _flush_pending()
            bullet_text = line.lstrip()[2:].strip()
            if current_section is not None and bullet_text:
                current_section["blocks"].append({
                    "type": "bullet",
                    "text": bullet_text,
                })
            continue
        # Skill category
        m = _SKILL_CATEGORY_RE.match(line.strip())
        if m and current_section is not None:
            _flush_pending()
            current_section["blocks"].append({
                "type": "skill_category",
                "text": (m.group(1).strip(), m.group(2).strip()),
            })
            continue
        # Blank line — flush
        if not line.strip():
            _flush_pending()
            continue
        # Plain text line
        if current_section is None:
            pre_section_lines.append(line)
        else:
            pending_paragraph_lines.append(line)

    _flush_pending()
    return {
        "name": name,
        "tagline": tagline,
        "contact": contact,
        "sections": sections,
    }


def _parse_cover_letter_markdown(content: str) -> dict:
    """Parse the cover letter markdown into:
      {
        "title_line": str (the '# Cover Letter — ...' line),
        "blocks": [str, str, ...]  # paragraphs in order, blank-separated
      }
    """
    lines = content.splitlines()
    title_line = ""
    blocks: list[str] = []
    current: list[str] = []
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("# ") and not title_line:
            title_line = line[2:].strip()
            continue
        if not line.strip():
            if current:
                blocks.append(
                    " ".join(p.strip() for p in current if p.strip())
                )
                current = []
            continue
        current.append(line)
    if current:
        blocks.append(" ".join(p.strip() for p in current if p.strip()))
    return {"title_line": title_line, "blocks": [b for b in blocks if b]}


# --- File-naming helpers -----------------------------------------

_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]+")


def _split_first_last(name: str) -> tuple[str, str]:
    """'ALEX DOE' / 'Alex Doe' → ('Alex', 'Doe').
    If only one token, last name is empty string."""
    parts = [p for p in re.split(r"\s+", name.strip()) if p]
    if not parts:
        return ("", "")
    titled = [p[:1].upper() + p[1:].lower() for p in parts]
    if len(titled) == 1:
        return (titled[0], "")
    return (titled[0], "".join(titled[1:]))


def _company_token(employer: str) -> str:
    """Strip non-alphanumerics so the file name stays portable."""
    return _NON_ALNUM_RE.sub("", employer or "Company")


def _resume_filename(name: str, employer: str, today: str) -> str:
    first, last = _split_first_last(name)
    company = _company_token(employer)
    return f"{first}{last}Resume{company}{today}.docx"


def _cover_letter_filename(
    name: str, employer: str, today: str,
) -> str:
    first, last = _split_first_last(name)
    company = _company_token(employer)
    return f"{first}{last}CoverLetter{company}{today}.docx"


# --- Main renderer class ------------------------------------------

class ATSResumeRenderer:
    """Render markdown drafts into ATS-friendly .docx files.

    All formatting rules are pinned to module-level constants so
    tests can verify them without re-parsing the docx XML.
    """

    def render_resume(
        self,
        content: str,
        posting: dict,
        output_dir: Path,
        *,
        today: Optional[str] = None,
    ) -> Path:
        parsed = _parse_resume_markdown(content)
        if not parsed["name"]:
            raise ValueError(
                "resume markdown must start with '# {NAME}' line"
            )

        doc = Document()
        _set_section_margins(doc)
        # Defensive: ensure header/footer have no content
        for section in doc.sections:
            section.header.is_linked_to_previous = True
            section.footer.is_linked_to_previous = True

        # Name + tagline + contact
        _name_paragraph(doc, parsed["name"])
        if parsed.get("tagline"):
            _centered_body_paragraph(doc, parsed["tagline"])
        if parsed.get("contact"):
            _centered_body_paragraph(doc, parsed["contact"])

        # Sections
        for section in parsed["sections"]:
            _heading_paragraph(doc, section["heading"])
            for block in section["blocks"]:
                if block["type"] == "paragraph":
                    _body_paragraph(doc, block["text"])
                elif block["type"] == "bullet":
                    _bullet_paragraph(doc, block["text"])
                elif block["type"] == "role_header":
                    _role_header_paragraph(doc, block["text"])
                elif block["type"] == "skill_category":
                    cat, terms = block["text"]
                    _skill_category_paragraph(doc, cat, terms)

        today = today or datetime.now(timezone.utc).strftime("%Y%m%d")
        output_dir.mkdir(parents=True, exist_ok=True)
        out_path = output_dir / _resume_filename(
            parsed["name"], posting.get("employer") or "", today,
        )
        doc.save(str(out_path))
        return out_path

    def render_cover_letter(
        self,
        content: str,
        posting: dict,
        candidate_name: str,
        output_dir: Path,
        *,
        today: Optional[str] = None,
    ) -> Path:
        parsed = _parse_cover_letter_markdown(content)
        if not parsed["title_line"]:
            raise ValueError(
                "cover letter markdown must start with "
                "'# Cover Letter — ...' line"
            )

        doc = Document()
        _set_section_margins(doc)
        for section in doc.sections:
            section.header.is_linked_to_previous = True
            section.footer.is_linked_to_previous = True

        # Cover letters don't have a centered banner — body-aligned
        # paragraphs only. Each parsed block becomes one paragraph.
        for block in parsed["blocks"]:
            _body_paragraph(doc, block)

        today = today or datetime.now(timezone.utc).strftime("%Y%m%d")
        output_dir.mkdir(parents=True, exist_ok=True)
        out_path = output_dir / _cover_letter_filename(
            candidate_name, posting.get("employer") or "", today,
        )
        doc.save(str(out_path))
        return out_path
